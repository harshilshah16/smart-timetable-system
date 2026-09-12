#!/usr/bin/env python3
import io
import os
import re
import uuid
import csv
from collections import defaultdict
from datetime import datetime, timezone

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


app = FastAPI(title="STM Timetable API", version="1.0.0")
app.add_middleware(
	CORSMiddleware,
	allow_origins=["*"],
	allow_credentials=True,
	allow_methods=["*"],
	allow_headers=["*"],
)

DEFAULT_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
DEFAULT_PERIODS = [
	("08:00", "09:00"),
	("09:00", "09:55"),
	("10:05", "11:00"),
	("11:00", "12:00"),
	("12:30", "13:30"),
	("13:30", "14:30"),
	("14:30", "15:30"),
	("15:30", "16:30"),
]
DIVISIONS = {"SY A", "SY B", "TY A", "TY B", "LY A", "LY B"}


def clean(value):
	return re.sub(r"\s+", " ", str(value or "").strip())


def key(value):
	return re.sub(r"[^a-z0-9]", "", clean(value).lower())


def first_value(row, names, default=""):
	values = {key(name): value for name, value in row.items()}
	for name in names:
		value = values.get(key(name), "")
		if clean(value):
			return clean(value)
	return default


def integer(value, default):
	try:
		return max(1, int(float(str(value).strip())))
	except (TypeError, ValueError):
		return default


def normalize_division(value):
	value = clean(value).upper().replace("SECTION", "").strip()
	value = re.sub(r"\s+", " ", value)
	aliases = {"SY-A": "SY A", "SY-B": "SY B", "TY-A": "TY A", "TY-B": "TY B", "LY-A": "LY A", "LY-B": "LY B"}
	return aliases.get(value, value)


def read_settings(workbook):
	settings = {}
	if "Settings" in workbook.sheetnames:
		for row in workbook["Settings"].iter_rows(values_only=True):
			if len(row) >= 2 and clean(row[0]):
				settings[key(row[0])] = clean(row[1])
	days = [clean(item) for item in settings.get("days", "Monday,Tuesday,Wednesday,Thursday,Friday").split(",") if clean(item)]
	period_text = settings.get("periods", "")
	periods = []
	for item in period_text.split(","):
		parts = [clean(part) for part in item.split("-")]
		if len(parts) == 2 and parts[0] and parts[1]:
			periods.append((parts[0], parts[1]))
	return days or DEFAULT_DAYS, periods or DEFAULT_PERIODS, settings


def read_courses(workbook):
	sheet_name = "Courses" if "Courses" in workbook.sheetnames else workbook.sheetnames[0]
	sheet = workbook[sheet_name]
	rows = list(sheet.iter_rows(values_only=True))
	if not rows:
		raise ValueError("The workbook has no course rows")
	headers = [clean(value) for value in rows[0]]
	required = {"division", "coursecode", "facultyid"}
	available = {key(value) for value in headers}
	if not required.issubset(available):
		raise ValueError("Courses sheet must include Division, Course Code, and Faculty ID columns")
	courses = []
	for row_number, values in enumerate(rows[1:], start=2):
		if not any(clean(value) for value in values):
			continue
		raw = dict(zip(headers, values))
		division = normalize_division(first_value(raw, ["Division", "Class", "Section"]))
		if division not in DIVISIONS:
			raise ValueError(f"Row {row_number}: Division must be one of {', '.join(sorted(DIVISIONS))}")
		course_type = first_value(raw, ["Course Type", "Type"], "Lecture").lower()
		room_value = first_value(raw, ["Room", "Room Preference", "Room Preferences"], "Any")
		lab_default = "lab" in course_type or room_value.lower() in {"lab", "labs", "any lab"}
		courses.append({
			"id": str(uuid.uuid4()),
			"division": division,
			"course_code": first_value(raw, ["Course Code", "Subject Code", "Code"]),
			"course_name": first_value(raw, ["Course Name", "Subject", "Course"], "Untitled course"),
			"short_name": first_value(raw, ["Short Name", "Subject Short Name", "Abbreviation"]),
			"faculty_id": first_value(raw, ["Faculty ID", "Instructor ID", "Faculty", "Instructor"]),
			"room": room_value,
			"sessions_per_week": integer(first_value(raw, ["Sessions Per Week", "Weekly Sessions", "Sessions"]), 1),
			"duration": integer(first_value(raw, ["Duration", "Duration (Periods)"], ""), 2 if lab_default else 1),
			"preferred_day": first_value(raw, ["Preferred Day", "Day"]),
			"preferred_period": first_value(raw, ["Preferred Period", "Period"]),
			"room_options": [clean(room) for room in room_value.split(",") if clean(room)],
			"course_type": course_type,
			"batch": first_value(raw, ["Batch", "Lab Batch", "Batch Name"]),
			"lab_group": first_value(raw, ["Lab Group", "Lab Subject", "Course Code", "Subject Code", "Code"]),
		})
	if not courses:
		raise ValueError("The workbook has no usable course rows")
	return courses


def read_lab_rooms(workbook):
	if "Rooms" not in workbook.sheetnames:
		return []
	sheet = workbook["Rooms"]
	rows = list(sheet.iter_rows(values_only=True))
	if not rows:
		return []
	headers = [clean(value) for value in rows[0]]
	rooms = []
	for row_number, values in enumerate(rows[1:], start=2):
		if not any(clean(value) for value in values):
			continue
		raw = dict(zip(headers, values))
		room_code = first_value(raw, ["Room No", "Room Number", "Lab Code", "Room Code", "Room", "Code"])
		room_type = first_value(raw, ["Type", "Room Type"], "Lab").lower()
		if room_code and "lab" in room_type:
			rooms.append({"code": room_code})
	return rooms


def overlaps(slots, day, period_index, duration):
	return any(slot["day"] == day and period_index < slot["period_index"] + slot["duration"] and slot["period_index"] < period_index + duration for slot in slots)


def display_subject(course):
	if course.get("short_name"):
		return course["short_name"]
	normalized = re.sub(r"[^a-z0-9]", "", course["course_name"].lower())
	known_names = {
		"datastructureslab": "DS",
		"datastructures": "DS",
		"dbmslab": "DBMS",
		"databasemanagementsystemlab": "DBMS",
		"dlcalab": "DLCA",
		"digitallogicandcomputerarchitecturelab": "DLCA",
		"oopmlab": "OOPM",
		"objectorientedprogramminglab": "OOPM",
		"applicationsofmathematics": "AME-I",
	}
	return known_names.get(normalized, course["course_name"])


def generate(courses, days, periods, labs):
	scheduled = []
	conflicts = []
	occupied_division = defaultdict(list)
	occupied_faculty = defaultdict(list)
	occupied_room = defaultdict(list)
	last_subject_period = defaultdict(list)
	day_load = defaultdict(int)
	division_day_load = defaultdict(int)
	lecture_day_load = defaultdict(int)
	division_lecture_day_load = defaultdict(int)

	def add_entry(course, session, candidate, duration, room, batch=""):
		item = {
			"id": course["id"], "division": course["division"], "course_code": course["course_code"],
			"course_name": course["course_name"], "faculty_id": course["faculty_id"], "room": room,
			"batch": batch, "session": session, "day": candidate["day"], "period_index": candidate["period_index"],
			"start_time": periods[candidate["period_index"]][0],
			"end_time": periods[candidate["period_index"] + duration - 1][1], "duration": duration,
		}
		scheduled.append(item)
		division_key = f'{course["division"]}|{batch}' if batch else course["division"]
		occupied_division[division_key].append(candidate)
		if course["faculty_id"]:
			occupied_faculty[course["faculty_id"]].append(candidate)
		if room.lower() != "any":
			occupied_room[room.lower()].append(candidate)
		day_load[candidate["day"]] += duration
		division_day_load[(course["division"], candidate["day"])] += duration
		if not batch:
			lecture_day_load[candidate["day"]] += duration
			division_lecture_day_load[(course["division"], candidate["day"])] += duration
		return item

	def free_for(course, candidate, room, batch=""):
		division_key = f'{course["division"]}|{batch}' if batch else course["division"]
		division_busy = occupied_division[division_key]
		if batch:
			division_busy = division_busy + occupied_division[course["division"]]
		else:
			for key_name, slots in occupied_division.items():
				if key_name.startswith(f'{course["division"]}|'):
					division_busy = division_busy + slots
		return not overlaps(division_busy, **candidate) and not overlaps(occupied_faculty[course["faculty_id"]], **candidate) and (room.lower() == "any" or not overlaps(occupied_room[room.lower()], **candidate))

	def room_choices(course):
		options = course["room_options"] or ["Any"]
		if len(options) == 1 and options[0].lower() in {"lab", "labs", "any lab"}:
			return [room["code"] for room in labs]
		return options

	def add_conflict(course, session, reason):
		conflicts.append({"course_code": course["course_code"], "course_name": course["course_name"], "division": course["division"], "faculty_id": course["faculty_id"], "batch": course.get("batch", ""), "session": session, "reason": reason})

	# Lab rotation: on each of four lab days, each batch receives a different lab subject.
	lab_courses = [course for course in courses if "lab" in course["course_type"] or any(option.lower() in {"lab", "labs", "any lab"} for option in course["room_options"])]
	lab_groups = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
	for course in lab_courses:
		lab_groups[course["division"]][course["lab_group"]][course["batch"] or "B1"].append(course)
	for division, groups in lab_groups.items():
		group_names = sorted(groups)
		batches = sorted({batch for group in groups.values() for batch in group} | {"B1", "B2", "B3", "B4"})
		if len(labs) < len(batches):
			for group_name in group_names:
				add_conflict(groups[group_name][batches[0]][0], 1, f"Rooms sheet needs at least {len(batches)} lab rooms for the lab rotation")
			continue
		lab_starts = [index for index in range(len(periods) - 1) if periods[index][1] == periods[index + 1][0]]
		for day_index, day in enumerate(days[:min(4, len(days))]):
			for batch_index, batch in enumerate(batches[:len(labs)]):
				group_name = group_names[(batch_index + day_index) % len(group_names)]
				course = groups[group_name].get(batch, groups[group_name].get("B1", []))
				if not course:
					add_conflict(next(iter(groups[group_name].values()))[0], day_index + 1, f"No Excel row for batch {batch} in lab group {group_name}")
					continue
				course = course[0]
				placed = False
				for period_index in lab_starts:
					candidate = {"day": day, "period_index": period_index, "duration": 2}
					room = labs[batch_index]["code"]
					if free_for(course, candidate, room, batch):
						add_entry(course, day_index + 1, candidate, 2, room, batch)
						last_subject_period[(division, group_name, batch)].append((day, period_index))
						placed = True
						break
				if not placed:
					add_conflict(course, day_index + 1, "No free two-period lab block satisfies batch, faculty, and room constraints")

	# Lectures and non-rotating courses use the remaining timetable slots.
	normal_courses = [course for course in courses if course not in lab_courses]
	units = [(course, session) for course in normal_courses for session in range(1, course["sessions_per_week"] + 1)]
	units.sort(key=lambda item: (-item[0]["duration"], item[0]["division"], item[0]["course_code"]))
	for course, session in units:
		duration = min(course["duration"], len(periods))
		best = None
		for day in days:
			if course["preferred_day"] and course["preferred_day"].lower() != day.lower():
				continue
			for period_index in range(len(periods) - duration + 1):
				candidate = {"day": day, "period_index": period_index, "duration": duration}
				if course["preferred_period"] and clean(course["preferred_period"]) not in {periods[period_index][0], f"{periods[period_index][0]}-{periods[period_index + duration - 1][1]}"}:
					continue
				for room in room_choices(course):
					if free_for(course, candidate, room) and not any(previous_day == day and abs(previous_period - period_index) <= duration for previous_day, previous_period in last_subject_period[(course["division"], course["course_code"])]):
						score = (division_lecture_day_load[(course["division"], day)], lecture_day_load[day], period_index)
						if best is None or score < best[0]:
							best = (score, candidate, room, day)
		if best is None:
			add_conflict(course, session, "No free day and period satisfies the division, faculty, and room constraints")
		else:
			_, candidate, room, day = best
			add_entry(course, session, candidate, duration, room)
			last_subject_period[(course["division"], course["course_code"])].append((day, candidate["period_index"]))
	return scheduled, conflicts


def build_result(workbook):
	days, periods, settings = read_settings(workbook)
	courses = read_courses(workbook)
	labs = read_lab_rooms(workbook)
	scheduled, conflicts = generate(courses, days, periods, labs)
	by_division = {division: [item for item in scheduled if item["division"] == division] for division in sorted(DIVISIONS)}
	return {
		"generated_at": datetime.now(timezone.utc).isoformat(),
		"days": days,
		"periods": [{"start": start, "end": end} for start, end in periods],
		"settings": settings,
		"labs": labs,
		"divisions": by_division,
		"entries": scheduled,
		"conflicts": conflicts,
		"summary": {"courses": len(courses), "scheduled_sessions": len(scheduled), "unscheduled_sessions": len(conflicts)},
	}


def sample_workbook():
	workbook = Workbook()
	courses = workbook.active
	courses.title = "Courses"
	courses.append(["Division", "Batch", "Course Type", "Course Code", "Course Name", "Faculty ID", "Room", "Sessions Per Week", "Duration", "Preferred Day", "Preferred Period"])
	lecture_rows = [
		["SY A", "", "Lecture", "C301", "Applications of Mathematics", "F001", "CR-10", 3, 1, "", ""],
		["SY A", "", "Lecture", "C302", "Data Structures", "F002", "CR-10", 3, 1, "", ""],
		["SY A", "", "Lecture", "C303", "Digital Logic and Computer Architecture", "F003", "CR-10", 3, 1, "", ""],
	]
	for row in lecture_rows:
		courses.append(row)
	for lab_code, lab_name, faculty_prefix in [("L301", "Data Structures Lab", "DS"), ("L302", "DBMS Lab", "DBMS"), ("L303", "DLCA Lab", "DLCA"), ("L304", "OOPM Lab", "OOPM")]:
		for batch in ["B1", "B2", "B3", "B4"]:
			courses.append(["SY A", batch, "Lab", lab_code, lab_name, f"{faculty_prefix}-{batch}", "Lab", 1, 2, "", ""])
	rooms = workbook.create_sheet("Rooms")
	rooms.append(["Room No", "Type"])
	for room in ["502A", "502B", "503A", "503B"]:
		rooms.append([room, "Lab"])
	settings = workbook.create_sheet("Settings")
	settings.append(["Setting", "Value"])
	settings.append(["College", "K J Somaiya Institute of Technology"])
	settings.append(["Department", "Department of Computer Engineering"])
	settings.append(["Academic Year", "2026-27"])
	settings.append(["Semester", "Odd Semester"])
	settings.append(["Days", "Monday,Tuesday,Wednesday,Thursday,Friday"])
	settings.append(["Periods", "08:00-09:00,09:00-09:55,10:05-11:00,11:00-12:00,12:30-13:30,13:30-14:30,14:30-15:30,15:30-16:30"])
	instructions = workbook.create_sheet("Instructions")
	instructions.append(["How to use this template"])
	instructions.append(["Keep one Courses row for every division, subject, batch, and faculty assignment."])
	instructions.append(["For labs, use Course Type = Lab, Room = Lab, Duration = 2, and Batch = B1/B2/B3/B4."])
	instructions.append(["The four lab subjects rotate across B1-B4 over four days so every batch completes every lab."])
	instructions.append(["Enter every real lab room number in Rooms. Do not use locations; room number uniquely identifies a lab."])
	instructions.append(["Use the same Faculty ID wherever a faculty member teaches another division; the generator will report clashes."])
	return workbook


@app.get("/health")
def health():
	return {"status": "ok", "service": "timetable-generator"}


@app.get("/api/timetable/sample.xlsx")
def download_sample_workbook():
	output = io.BytesIO()
	sample_workbook().save(output)
	output.seek(0)
	return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=timetable-sample.xlsx"})


@app.post("/api/timetable/generate")
async def generate_timetable(file: UploadFile = File(...)):
	if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm", ".csv")):
		raise HTTPException(status_code=400, detail="Upload an .xlsx, .xlsm, or .csv file")
	try:
		contents = await file.read()
		if file.filename.lower().endswith(".csv"):
			workbook = Workbook()
			sheet = workbook.active
			for row in csv.reader(io.StringIO(contents.decode("utf-8-sig"))):
				sheet.append(row)
		else:
			workbook = load_workbook(io.BytesIO(contents), data_only=True)
		return build_result(workbook)
	except ValueError as error:
		raise HTTPException(status_code=422, detail=str(error)) from error
	except Exception as error:
		raise HTTPException(status_code=400, detail=f"Could not read workbook: {error}") from error


@app.post("/api/timetable/generate.xlsx")
async def generate_timetable_xlsx(file: UploadFile = File(...)):
	result = await generate_timetable(file)
	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Timetable"
	headers = ["Division", "Course Code", "Course Name", "Faculty ID", "Room", "Day", "Start", "End", "Session"]
	fields = ["division", "course_code", "course_name", "faculty_id", "room", "day", "start_time", "end_time", "session"]
	sheet.append(headers)
	for entry in result["entries"]:
		sheet.append([entry[field] for field in fields])
	output = io.BytesIO()
	workbook.save(output)
	output.seek(0)
	return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=timetable.xlsx"})


@app.post("/api/timetable/generate.pdf")
async def generate_timetable_pdf(file: UploadFile = File(...)):
	result = await generate_timetable(file)
	output = io.BytesIO()
	document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=7 * mm, leftMargin=7 * mm, topMargin=7 * mm, bottomMargin=7 * mm)
	styles = getSampleStyleSheet()
	settings = result.get("settings", {})
	college = settings.get("college", "K J Somaiya Institute of Technology")
	department = settings.get("department", "Department of Computer Engineering")
	academic_year = settings.get("academic_year", "2026-27")
	semester = settings.get("semester", "Odd Semester")
	periods = result["periods"]
	break_after = {1, 3} if len(periods) >= 4 else set()
	flowables = []

	def cell_text(value, size=6.5, bold=False):
		font = "Helvetica-Bold" if bold else "Helvetica"
		text = clean(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
		return Paragraph(f'<font name="{font}" size="{size}">{text}</font>', styles["Normal"])

	def subject_color(course_code):
		palette = ["#b7d2f0", "#f3a6b8", "#b8d49b", "#d6a6c1", "#9ddfe5", "#f5e88a", "#c7b5e8"]
		return colors.HexColor(palette[sum(ord(char) for char in course_code) % len(palette)])

	for division_index, (division, entries) in enumerate(result["divisions"].items()):
		if division_index:
			flowables.append(PageBreak())
		flowables.extend([
			Paragraph(f'<b>{college}</b>', styles["Title"]),
			Paragraph(department, styles["Normal"]),
			Spacer(1, 1.5 * mm),
			Paragraph(f'<b>TIME TABLE FOR ACADEMIC YEAR : {academic_year} ({semester})</b>', styles["Heading2"]),
			Paragraph(f'<b>Class : {division}</b>', styles["Normal"]),
			Spacer(1, 3 * mm),
		])
		columns = ["Day/Time"]
		column_types = ["day"]
		for index, period in enumerate(periods):
			columns.append(f'{period["start"]}\n{period["end"]}')
			column_types.append("period")
			if index in break_after:
				columns.append("BREAK")
				column_types.append("break")
		rows = [[cell_text(value, 6.5, True) for value in columns]]
		style_commands = [
			("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e2e2")),
			("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#4d4d4d")),
			("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
			("ALIGN", (1, 0), (-1, -1), "CENTER"),
			("ALIGN", (0, 0), (0, -1), "CENTER"),
		]
		entry_by_day = defaultdict(list)
		for entry in entries:
			entry_by_day[entry["day"]].append(entry)
		for day in result["days"]:
			row = [cell_text(day, 7, True)]
			period_column = 1
			for period_index in range(len(periods)):
				matches = [entry for entry in entry_by_day[day] if entry["period_index"] == period_index]
				if matches:
					lines = []
					for match in sorted(matches, key=lambda entry: entry["batch"] or ""):
						if match["batch"]:
							lines.append(f'{match["batch"]} {display_subject(match)} / {match["room"]}')
						else:
							lines.append(f'{display_subject(match)} / {match["faculty_id"]} / {match["room"]}')
					row.append(cell_text("\n".join(lines), 6 if len(matches) > 1 else 5.8))
					for match in matches:
						style_commands.append(("BACKGROUND", (period_column, len(rows)), (period_column + match["duration"] - 1, len(rows)), subject_color(match["course_code"])))
						if match["duration"] > 1 and period_index + match["duration"] <= len(periods):
							span_end = period_column + match["duration"] - 1
							if all(column_types[index + 1] == "period" for index in range(period_index, period_index + match["duration"])):
								style_commands.append(("SPAN", (period_column, len(rows)), (span_end, len(rows))))
				else:
					row.append("")
				period_column += 1
				if period_index in break_after:
					row.append(cell_text("BREAK", 6, True))
					style_commands.append(("BACKGROUND", (period_column, len(rows)), (period_column, len(rows)), colors.yellow))
					period_column += 1
			rows.append(row)
		table = Table(rows, colWidths=[25 * mm] + [((277 - 25 - 10) / len(columns)) * mm] * (len(columns) - 1), rowHeights=[13 * mm] + [22 * mm] * len(result["days"]), repeatRows=1)
		table.setStyle(TableStyle(style_commands))
		flowables.extend([table, Spacer(1, 4 * mm), Paragraph("Course Code | Course Name | Faculty | Room", styles["Normal"])])
		legend_rows = [[cell_text("Course Code", 6.5, True), cell_text("Course Name", 6.5, True), cell_text("Faculty", 6.5, True), cell_text("Room", 6.5, True)]]
		seen = set()
		for entry in entries:
			if entry["course_code"] in seen:
				continue
			seen.add(entry["course_code"])
			legend_rows.append([cell_text(entry["course_code"]), cell_text(entry["course_name"]), cell_text(entry["faculty_id"]), cell_text(entry["room"])])
		legend = Table(legend_rows, colWidths=[25 * mm, 105 * mm, 45 * mm, 35 * mm], repeatRows=1)
		legend.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e2e2")), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#707070")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
		flowables.extend([legend, Spacer(1, 2 * mm), Paragraph(f'Generated from uploaded Excel data on {result["generated_at"]}', styles["Normal"])])
	if result["conflicts"]:
		flowables.extend([PageBreak(), Paragraph("Unscheduled Sessions", styles["Title"])])
		conflict_rows = [[cell_text(header, 7, True) for header in ["Division", "Course", "Faculty", "Session", "Reason"]]]
		for conflict in result["conflicts"]:
			conflict_rows.append([cell_text(conflict["division"]), cell_text(f'{conflict["course_code"]} - {conflict["course_name"]}'), cell_text(conflict["faculty_id"]), cell_text(conflict["session"]), cell_text(conflict["reason"])])
		conflict_table = Table(conflict_rows, repeatRows=1, colWidths=[25 * mm, 75 * mm, 25 * mm, 18 * mm, 110 * mm])
		conflict_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#9b2c2c")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d6b5b5")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
		flowables.append(conflict_table)
	document.build(flowables)
	output.seek(0)
	return StreamingResponse(output, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=timetable.pdf"})


if __name__ == "__main__":
	import uvicorn
	uvicorn.run(app, host=os.getenv("HOST", "0.0.0.0"), port=int(os.getenv("PORT", "8000")))
